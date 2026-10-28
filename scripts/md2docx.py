import sys
import io
import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BODY_CJK = "微软雅黑"
BODY_LATIN = "Segoe UI"
MONO = "Consolas"
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x55, 0x55)
GREEN = RGBColor(0x1B, 0x7A, 0x3D)
RED = RGBColor(0xB4, 0x23, 0x18)


def set_run_font(run, cjk=BODY_CJK, latin=BODY_LATIN, size=None, bold=None, color=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)


def style_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_LATIN
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CJK)
    pf = normal.paragraph_format
    pf.line_spacing = 1.45
    pf.space_after = Pt(6)


def shade(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, text in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(text)
        set_run_font(run, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(hdr[i], "2F4F6F")
    for row in body:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            if i >= len(cells):
                break
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(text)
            set_run_font(run, size=9.5)
    doc.add_paragraph()
    return table


INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text, base_size=10.5, base_color=None):
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            set_run_font(run, size=base_size, color=base_color)
        token = m.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, cjk=MONO, latin=MONO, size=base_size - 0.5, color=RGBColor(0xA0, 0x30, 0x30))
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, bold=True, color=base_color)
        else:
            label = token[1:token.index("]")]
            run = paragraph.add_run(label)
            set_run_font(run, size=base_size, color=RGBColor(0x1B, 0x5E, 0xA0))
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=base_size, color=base_color)


def convert(md_path, out_path):
    lines = io.open(md_path, encoding="utf-8").read().split("\n")
    doc = Document()
    style_document(doc)
    sec = doc.sections[0]
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.4)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.6)

    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run("\n".join(code_lines))
                set_run_font(run, cjk=MONO, latin=MONO, size=9, color=RGBColor(0x24, 0x29, 0x2E))
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = []
            for raw in block:
                cells = [c.strip() for c in raw.strip("|").split("|")]
                if all(set(c) <= set("-: ") and c for c in cells):
                    continue
                rows.append(cells)
            if rows:
                add_table(doc, rows)
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(text)
                set_run_font(run, size=20, bold=True, color=RGBColor(0x12, 0x2A, 0x43))
                sub = doc.add_paragraph()
                sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
                sub.paragraph_format.space_after = Pt(14)
                run = sub.add_run("CordiBank　面向金融的蜂群多智能体系统")
                set_run_font(run, size=9.5, color=MUTED)
            else:
                size = {2: 15, 3: 13, 4: 11.5, 5: 11, 6: 10.5}.get(level, 11)
                color = RGBColor(0x12, 0x2A, 0x43) if level <= 3 else RGBColor(0x2B, 0x3A, 0x4A)
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12 if level <= 3 else 8)
                p.paragraph_format.space_after = Pt(5)
                run = p.add_run(text)
                set_run_font(run, size=size, bold=True, color=color)
                if level == 2:
                    ppr = p._element.get_or_add_pPr()
                    pbdr = OxmlElement("w:pBdr")
                    bottom = OxmlElement("w:bottom")
                    bottom.set(qn("w:val"), "single")
                    bottom.set(qn("w:sz"), "8")
                    bottom.set(qn("w:space"), "4")
                    bottom.set(qn("w:color"), "2F4F6F")
                    pbdr.append(bottom)
                    ppr.append(pbdr)
            i += 1
            continue

        m = re.match(r"^\s*[-*]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, m.group(1))
            i += 1
            continue

        m = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, m.group(2))
            i += 1
            continue

        m = re.match(r"^>\s?(.*)$", stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            run = p.add_run(m.group(1))
            set_run_font(run, size=10.5, color=RGBColor(0x2B, 0x3A, 0x4A))
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, stripped)
        i += 1

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    src = sys.argv[1]
    dst = sys.argv[2]
    convert(src, dst)
    print("written", dst, os.path.getsize(dst))
